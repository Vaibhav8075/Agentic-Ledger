from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.dml import MSO_THEME_COLOR
from pptx.enum.text import MSO_AUTO_SIZE

OUT = "Paytm_Vyapar_Saathi_Pitch.pptx"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

NAVY = RGBColor(9, 20, 38)
INK = RGBColor(21, 31, 45)
MUTED = RGBColor(99, 112, 128)
WHITE = RGBColor(250, 252, 255)
CYAN = RGBColor(0, 180, 210)
TEAL = RGBColor(0, 145, 137)
ORANGE = RGBColor(245, 143, 61)
RED = RGBColor(222, 79, 81)
PALE = RGBColor(235, 245, 248)
LIGHT = RGBColor(246, 248, 251)
LINE = RGBColor(215, 224, 232)
GREEN = RGBColor(42, 157, 112)

FONT = "Aptos"
DISPLAY = "Aptos Display"


def add_text(slide, text, x, y, w, h, size=18, color=INK, bold=False,
             align=PP_ALIGN.LEFT, font=FONT, valign=MSO_ANCHOR.TOP,
             margin=0.04, italic=False):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin)
    tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    return box


def add_rich_text(slide, runs, x, y, w, h, size=18, color=INK, align=PP_ALIGN.LEFT,
                  valign=MSO_ANCHOR.TOP, margin=0.04):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin)
    tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    for text, bold, color in runs:
        run = p.add_run()
        run.text = text
        run.font.name = FONT
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
    return box


def rect(slide, x, y, w, h, fill, line=None, radius=False):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line or fill
    if radius:
        shape.adjustments[0] = 0.08
    return shape


def line(slide, x1, y1, x2, y2, color=LINE, width=1.4, dash=None):
    shape = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    shape.line.color.rgb = color
    shape.line.width = Pt(width)
    if dash:
        shape.line.dash_style = dash
    return shape


def circle(slide, x, y, d, fill, line_color=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line_color or fill
    return shape


def title(slide, kicker, heading, sub=None, dark=False):
    color = WHITE if dark else INK
    muted = RGBColor(180, 198, 211) if dark else MUTED
    add_text(slide, kicker.upper(), 0.7, 0.42, 4.8, 0.28, 10, CYAN, True)
    add_text(slide, heading, 0.7, 0.78, 11.8, 0.62, 27, color, True, font=DISPLAY)
    if sub:
        add_text(slide, sub, 0.72, 1.48, 11.5, 0.42, 12.5, muted)


def footer(slide, n, dark=False):
    color = RGBColor(152, 171, 188) if dark else MUTED
    add_text(slide, "PAYTM VYAPAR SAATHI  /  AGENT LABS", 0.7, 7.12, 5.5, 0.18, 8.5, color, True)
    add_text(slide, f"{n:02d}", 12.1, 7.08, 0.45, 0.22, 9, color, True, align=PP_ALIGN.RIGHT)


def add_bullet(slide, text, x, y, w, color=INK, accent=CYAN, size=14, h=0.38):
    circle(slide, x, y + 0.07, 0.11, accent)
    add_text(slide, text, x + 0.22, y, w - 0.22, h, size, color)


def add_metric(slide, value, label, x, y, w, accent=CYAN, dark=False):
    fill = RGBColor(17, 35, 58) if dark else WHITE
    rect(slide, x, y, w, 1.08, fill, fill, True)
    add_text(slide, value, x + 0.16, y + 0.17, w - 0.32, 0.38, 22, accent, True, font=DISPLAY)
    add_text(slide, label, x + 0.16, y + 0.62, w - 0.32, 0.25, 10, RGBColor(184, 199, 211) if dark else MUTED, True)


def new_slide(bg=WHITE):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    rect(slide, 0, 0, 13.333, 7.5, bg)
    return slide


# Slide 1
slide = new_slide(NAVY)
rect(slide, 0, 0, 0.22, 7.5, CYAN)
add_text(slide, "PAYTM AGENT LABS  /  HACKATHON PITCH", 0.82, 0.62, 6.8, 0.3, 11, CYAN, True)
add_text(slide, "Vyapar Saathi", 0.78, 1.42, 8.2, 0.78, 42, WHITE, True, font=DISPLAY)
add_text(slide, "From payment anxiety to intelligent cashflow.", 0.82, 2.35, 8.8, 0.48, 20, RGBColor(207, 223, 232), False)
line(slide, 0.82, 3.2, 4.0, 3.2, CYAN, 2.5)
add_text(slide, "An autonomous multi-agent system for India's everyday merchant.", 0.82, 3.48, 6.5, 0.55, 15, RGBColor(177, 198, 211))
# right-side signal composition
circle(slide, 9.0, 1.15, 2.55, RGBColor(14, 45, 68), RGBColor(14, 45, 68))
circle(slide, 9.55, 1.70, 1.45, CYAN, CYAN)
add_text(slide, "₹", 9.55, 1.93, 1.45, 0.65, 38, NAVY, True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE, font=DISPLAY)
for yy, label, value, accent in [(4.5, "PAYMENT TRUST", "< 600 ms SLA", CYAN), (5.35, "CASHFLOW SAFETY", "EMI first", ORANGE), (6.2, "DEPLOYMENT", "Dockerised", TEAL)]:
    add_text(slide, label, 8.3, yy, 2.4, 0.2, 9, RGBColor(142, 168, 184), True)
    add_text(slide, value, 10.75, yy - 0.04, 1.7, 0.3, 14, accent, True, align=PP_ALIGN.RIGHT)
footer(slide, 1, True)

# Slide 2
slide = new_slide(LIGHT)
title(slide, "01  /  The merchant reality", "Two moments. One fragile trust loop.", "The happy path works. Real merchants live in the exceptions.")
# left story card
rect(slide, 0.72, 2.18, 5.85, 3.8, WHITE, WHITE, True)
add_text(slide, "01", 1.05, 2.52, 0.55, 0.4, 23, CYAN, True, font=DISPLAY)
add_text(slide, "The silent Soundbox", 1.72, 2.55, 3.9, 0.35, 17, INK, True)
add_text(slide, "A genuine payment lands during a 2G drop.\nThe screen says paid. The device says nothing.", 1.05, 3.05, 4.95, 0.72, 15, MUTED)
line(slide, 1.05, 4.07, 6.15, 4.07, LINE, 1)
add_text(slide, "MERCHANT COST", 1.05, 4.37, 1.4, 0.2, 9, MUTED, True)
add_text(slide, "Checkout stops. Anxiety rises. A real payment looks like a spoof.", 1.05, 4.66, 4.95, 0.56, 14, RED, True)
# right story card
rect(slide, 6.82, 2.18, 5.8, 3.8, NAVY, NAVY, True)
add_text(slide, "02", 7.15, 2.52, 0.55, 0.4, 23, ORANGE, True, font=DISPLAY)
add_text(slide, "The 4:00 PM crunch", 7.82, 2.55, 3.9, 0.35, 17, WHITE, True)
add_text(slide, "The distributor needs ₹4,500 now.\nThe merchant has ₹6,000 unsettled.", 7.15, 3.05, 4.95, 0.72, 15, RGBColor(184, 201, 211))
line(slide, 7.15, 4.07, 12.25, 4.07, RGBColor(49, 71, 91), 1)
add_text(slide, "MERCHANT COST", 7.15, 4.37, 1.4, 0.2, 9, RGBColor(150, 174, 188), True)
add_text(slide, "A full sweep can break tomorrow's ₹3,000 loan EMI.", 7.15, 4.66, 4.95, 0.56, 14, ORANGE, True)
add_text(slide, "We built for the edge cases that decide whether a merchant trusts the rail.", 0.75, 6.45, 11.9, 0.35, 15, INK, True, align=PP_ALIGN.CENTER)
footer(slide, 2)

# Slide 3
slide = new_slide(WHITE)
title(slide, "02  /  Real-time trust", "Spoof Guard separates fraud from lag.", "The agent reads the transaction hash and the Soundbox telemetry together.")
# flow
steps = [("PAYMENT", "₹50 transaction", CYAN), ("GATEWAY", "Hash verified", TEAL), ("TELEMETRY", "MQTT / QoS 1", ORANGE), ("P4B", "Actionable push", GREEN)]
xs = [0.82, 3.7, 6.58, 9.46]
for i, ((label, sub, accent), x) in enumerate(zip(steps, xs)):
    rect(slide, x, 2.55, 2.2, 1.35, PALE if i != 2 else RGBColor(255, 245, 232), PALE, True)
    circle(slide, x + 0.18, 2.78, 0.36, accent)
    add_text(slide, str(i + 1), x + 0.18, 2.81, 0.36, 0.2, 9, WHITE, True, align=PP_ALIGN.CENTER)
    add_text(slide, label, x + 0.68, 2.72, 1.35, 0.22, 10, accent, True)
    add_text(slide, sub, x + 0.68, 3.05, 1.35, 0.28, 12, INK, True)
    if i < 3:
        line(slide, x + 2.2, 3.22, xs[i + 1], 3.22, LINE, 1.5)
# result panels
rect(slide, 0.82, 4.62, 5.65, 1.35, RGBColor(255, 242, 242), RGBColor(255, 226, 226), True)
add_text(slide, "NO VALID MATCH", 1.08, 4.88, 1.75, 0.2, 10, RED, True)
add_text(slide, "FRAUD ALERT", 1.08, 5.2, 2.6, 0.35, 20, RED, True, font=DISPLAY)
add_text(slide, "Fake or replayed receipt → P4B alert", 3.15, 5.22, 2.9, 0.3, 12, INK)
rect(slide, 6.84, 4.62, 5.65, 1.35, RGBColor(235, 249, 246), RGBColor(208, 235, 227), True)
add_text(slide, "HASH MATCH + PUBACK DELAY", 7.1, 4.88, 2.65, 0.2, 10, GREEN, True)
add_text(slide, "PAYMENT SAFE", 7.1, 5.2, 2.35, 0.35, 20, GREEN, True, font=DISPLAY)
add_text(slide, "Audio delayed due to network → anxiety removed", 9.45, 5.22, 2.55, 0.3, 12, INK)
add_metric(slide, "< 600 ms", "LATENCY BUDGET", 0.82, 6.28, 2.4, CYAN)
add_metric(slide, "15–45 sec", "2G BUFFERING WINDOW", 3.5, 6.28, 2.4, ORANGE)
add_metric(slide, "QoS 1", "MQTT SIGNAL", 6.18, 6.28, 2.4, TEAL)
add_metric(slide, "P4B", "NATIVE ACTION", 8.86, 6.28, 2.4, GREEN)
footer(slide, 3)

# Slide 4
slide = new_slide(NAVY)
title(slide, "03  /  Intelligent liquidity", "Forecast tonight, but trust only the floor.", "The Liquidity Engine protects inventory and preserves the merchant's repayment path.", dark=True)
# left scenario
rect(slide, 0.8, 2.2, 5.0, 3.85, RGBColor(17, 35, 58), RGBColor(36, 60, 82), True)
add_text(slide, "THE INPUT", 1.12, 2.55, 1.2, 0.2, 10, CYAN, True)
add_text(slide, "Amul invoice", 1.12, 2.92, 2.7, 0.35, 20, WHITE, True, font=DISPLAY)
add_text(slide, "Supplier asks for immediate payment", 1.12, 3.42, 3.8, 0.25, 13, RGBColor(184, 201, 211))
line(slide, 1.12, 4.05, 5.35, 4.05, RGBColor(44, 70, 91), 1)
for y, label, val, col in [(4.35, "Invoice due", "₹4,500", ORANGE), (4.78, "Unsettled balance", "₹6,000", RED), (5.21, "Tomorrow's EMI", "₹3,000", CYAN)]:
    add_text(slide, label, 1.12, y, 2.35, 0.2, 11, RGBColor(164, 186, 198))
    add_text(slide, val, 4.1, y - 0.04, 1.25, 0.27, 14, col, True, align=PP_ALIGN.RIGHT)
# right prediction
add_text(slide, "THE PREDICTION", 6.5, 2.55, 1.7, 0.2, 10, ORANGE, True)
add_text(slide, "Tonight's evening rush", 6.5, 2.92, 4.8, 0.35, 20, WHITE, True, font=DISPLAY)
add_text(slide, "PyTorch LSTM, scored on held-out days", 6.5, 3.42, 5.2, 0.25, 13, RGBColor(184, 201, 211))
# chart
line(slide, 6.6, 5.65, 12.18, 5.65, RGBColor(61, 86, 105), 1)
line(slide, 6.6, 4.12, 6.6, 5.65, RGBColor(61, 86, 105), 1)
points = [(6.85, 5.38), (7.65, 5.2), (8.45, 5.31), (9.25, 4.95), (10.05, 4.67), (10.85, 4.28), (11.65, 4.02)]
for (x1, y1), (x2, y2) in zip(points, points[1:]):
    line(slide, x1, y1, x2, y2, ORANGE, 3)
for x, y in points:
    circle(slide, x - 0.06, y - 0.06, 0.12, ORANGE)
add_text(slide, "₹6,894", 10.75, 3.83, 1.45, 0.45, 28, ORANGE, True, font=DISPLAY, align=PP_ALIGN.RIGHT)
add_text(slide, "floor · ₹9,061 expected", 10.25, 4.3, 1.95, 0.22, 10, RGBColor(167, 188, 199), align=PP_ALIGN.RIGHT)
add_text(slide, "The AI does not ask: ‘Can they pay now?’", 6.5, 6.08, 5.7, 0.28, 14, WHITE, True)
add_text(slide, "It asks: ‘Can they pay and still stay healthy tomorrow?’", 6.5, 6.42, 5.7, 0.28, 14, CYAN, True)
footer(slide, 4, True)

# Slide 5
slide = new_slide(WHITE)
title(slide, "04  /  Safety by construction", "The forecast can help. The EMI reserve is never touched.", "A deterministic plan decides; the agent re-checks it before anything is paid.")
# formula band
rect(slide, 0.78, 2.18, 11.8, 1.22, NAVY, NAVY, True)
add_text(slide, "SETTLE NOW  ≤  BALANCE − EMI      ADVANCE  ≤  MIN(FORECAST FLOOR, LIMIT)", 1.1, 2.48, 11.1, 0.45, 17, WHITE, True, align=PP_ALIGN.CENTER, font=DISPLAY)
# numbers
items = [("₹3,000", "settled from balance", ORANGE), ("₹1,500", "advance, repaid tonight", CYAN), ("₹3,000", "kept for the EMI", GREEN)]
for i, (value, label, col) in enumerate(items):
    x = 1.1 + i * 4.02
    rect(slide, x, 4.08, 3.45, 1.23, LIGHT, LIGHT, True)
    add_text(slide, value, x + 0.18, 4.35, 3.1, 0.3, 22, col, True, align=PP_ALIGN.CENTER, font=DISPLAY)
    add_text(slide, label.upper(), x + 0.18, 4.78, 3.1, 0.2, 9, MUTED, True, align=PP_ALIGN.CENTER)
    if i < 2:
        add_text(slide, "+" if i == 0 else "|", x + 3.55, 4.43, 0.45, 0.25, 20, MUTED, True, align=PP_ALIGN.CENTER)
# outcome
rect(slide, 0.78, 5.85, 5.7, 0.7, RGBColor(235, 249, 246), RGBColor(208, 235, 227), True)
add_text(slide, "APPROVED", 1.05, 6.08, 1.25, 0.22, 12, GREEN, True)
add_text(slide, "Inventory stays stocked", 2.55, 6.08, 2.8, 0.22, 13, INK, True)
rect(slide, 6.85, 5.85, 5.7, 0.7, RGBColor(255, 246, 234), RGBColor(247, 226, 193), True)
add_text(slide, "PROTECTED", 7.12, 6.08, 1.35, 0.22, 12, ORANGE, True)
add_text(slide, "Credit path stays intact", 8.72, 6.08, 2.8, 0.22, 13, INK, True)
footer(slide, 5)

# Slide 6
slide = new_slide(LIGHT)
title(slide, "05  /  How it works", "A supervisory agent coordinates three focused workers.", "Every external action passes through PII controls, state checks, and an auditable integration layer.")
# master
rect(slide, 4.35, 2.08, 4.65, 1.05, NAVY, NAVY, True)
add_text(slide, "PHINITE MASTER AGENT", 4.65, 2.34, 4.05, 0.22, 12, CYAN, True, align=PP_ALIGN.CENTER)
add_text(slide, "State machine  •  PII masking  •  Async routing", 4.65, 2.7, 4.05, 0.2, 10, WHITE, align=PP_ALIGN.CENTER)
# workers
workers = [(0.78, "SPOOF GUARD", "MQTT + hash check", CYAN), (4.35, "INVOICE PARSER", "WhatsApp / OCR", TEAL), (7.92, "LIQUIDITY ENGINE", "Forecast + settle", ORANGE)]
for x, label, sub, col in workers:
    rect(slide, x, 4.05, 3.15, 1.12, WHITE, LINE, True)
    circle(slide, x + 0.25, 4.34, 0.3, col)
    add_text(slide, label, x + 0.7, 4.28, 2.15, 0.25, 11, INK, True)
    add_text(slide, sub, x + 0.7, 4.66, 2.15, 0.2, 10, MUTED)
    line(slide, x + 1.58, 3.13, x + 1.58, 4.05, LINE, 1.4)
# guardrail and integration
line(slide, 2.35, 5.17, 10.98, 5.17, LINE, 1.4)
rect(slide, 3.3, 5.52, 6.72, 0.65, RGBColor(235, 245, 248), RGBColor(202, 228, 235), True)
add_text(slide, "DETERMINISTIC GUARDRAIL  •  maker-checker  •  no unsafe full sweep", 3.55, 5.73, 6.2, 0.22, 11, TEAL, True, align=PP_ALIGN.CENTER)
line(slide, 6.66, 6.17, 6.66, 6.5, LINE, 1.4)
add_text(slide, "PAYTM INTEGRATION LAYER", 3.8, 6.55, 5.75, 0.22, 11, NAVY, True, align=PP_ALIGN.CENTER)
add_text(slide, "Soundbox telemetry  •  P4B deep link  •  HMAC settlement  •  NPCI idempotency", 1.65, 6.82, 10.05, 0.2, 9.5, MUTED, align=PP_ALIGN.CENTER)
footer(slide, 6)

# Slide 7
slide = new_slide(NAVY)
title(slide, "06  /  Production readiness", "Built to move from demo to dependable infrastructure.", "A focused architecture for trust, liquidity, and safe execution.", dark=True)
features = [
    ("FASTAPI", "Idempotent APIs\nand 504 recovery", CYAN),
    ("PHINITE", "Maker-checker\n+ PII redaction", TEAL),
    ("DOCKER", "Containerized\nbackend + live UI", ORANGE),
    ("OBSERVABILITY", "Actionable states\nfrom lag to success", GREEN),
]
for i, (label, desc, col) in enumerate(features):
    x = 0.82 + (i % 2) * 6.0
    y = 2.25 + (i // 2) * 1.62
    rect(slide, x, y, 5.25, 1.18, RGBColor(17, 35, 58), RGBColor(40, 65, 85), True)
    rect(slide, x, y, 0.12, 1.18, col, col)
    add_text(slide, label, x + 0.38, y + 0.25, 1.65, 0.2, 11, col, True)
    add_text(slide, desc, x + 2.0, y + 0.2, 2.75, 0.5, 14, WHITE, True)
# closing statement
line(slide, 0.82, 5.85, 12.45, 5.85, RGBColor(47, 72, 91), 1)
add_text(slide, "Static ledgers become intelligent cashflow engines.", 0.82, 6.18, 8.9, 0.38, 23, WHITE, True, font=DISPLAY)
add_text(slide, "Paytm Vyapar Saathi", 9.35, 6.24, 3.1, 0.28, 14, CYAN, True, align=PP_ALIGN.RIGHT)
footer(slide, 7, True)

# Set document metadata
prs.core_properties.title = "Paytm Vyapar Saathi - Hackathon Pitch"
prs.core_properties.subject = "Autonomous multi-agent system for merchant payment trust and liquidity"
prs.core_properties.author = "Paytm Agent Labs"
prs.core_properties.keywords = "Paytm, Vyapar Saathi, Phinite, Spoof Guard, Liquidity Engine"

prs.save(OUT)
print(OUT)
